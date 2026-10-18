from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

wb = Workbook()
F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="0B3B33")
B = Font(name=F); BB = Font(name=F, bold=True)
BLUE = Font(name=F, color="0000FF"); YEL = PatternFill("solid", fgColor="FFFF00"); TINT = PatternFill("solid", fgColor="E1F5EE")
thin = Side(style="thin", color="BFBFBF"); BOR = Border(top=thin, bottom=thin, left=thin, right=thin)
RP = '"Rp"#,##0;("Rp"#,##0);-'
def hdr(ws, row, cols):
    for j, t in enumerate(cols, 1):
        c = ws.cell(row=row, column=j, value=t); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOR
def put(ws, r, c, v, font=B, fmt=None, fill=None):
    cell = ws.cell(row=r, column=c, value=v); cell.font = font; cell.border = BOR
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell

# ---------------- Sheet 1: Biaya dasar ----------------
ws = wb.active; ws.title = "Biaya dasar"
ws["A1"] = "Biaya dasar (estimasi bottom-up) per lokasi"; ws["A1"].font = Font(name=F, bold=True, size=13)
ws["A2"] = "Sel biru = angka masukan yang boleh diubah; sel hitam = rumus. Harga marketplace Indonesia, Oktober 2026."; ws["A2"].font = Font(name=F, italic=True, size=9)
hdr(ws, 4, ["No", "Komponen per lokasi", "Spesifikasi", "Jml", "Harga satuan (Rp)", "Subtotal (Rp)"])
items = [
 ("Water meter keluar multi-jet kuningan", "DN40 kelas B + modul pulsa 10 L", 1, 1600000),
 ("Water meter masuk multi-jet kuningan", "DN20/DN25 kelas B + modul pulsa 10 L", 1, 1050000),
 ("Strainer Y kuningan", "1,5 inci", 1, 200000),
 ("Sensor level tekanan hidrostatik", "0–2 mH2O, 4–20 mA, tee dan shunt", 1, 280000),
 ("Fitting, union, tee, seal tape, lem", "", 1, 200000),
 ("ESP32 DevKit V1", "", 1, 75000),
 ("Modul 4G SIM7600 dan antena", "", 1, 500000),
 ("Kartu SIM IM3", "paket 1 GB", 1, 50000),
 ("Modul RTC DS3231 dan microSD 8 GB", "", 1, 85000),
 ("PCB, resistor, kapasitor, terminal", "", 1, 50000),
 ("Adaptor 12 V, charger, aki 7 Ah, 2 step-down", "catu daya PLN + cadangan", 1, 390000),
 ("Box IP65, kabel, klem", "", 1, 220000),
 ("Jasa tukang pipa", "tarif rencana per lokasi", 1, 250000),
]
r = 5
for i, (n, s, q, h) in enumerate(items, 1):
    put(ws, r, 1, i); put(ws, r, 2, n); put(ws, r, 3, s)
    put(ws, r, 4, q, BLUE); put(ws, r, 5, h, BLUE, RP); put(ws, r, 6, f"=D{r}*E{r}", B, RP); r += 1
last = r - 1
put(ws, r, 2, "Per lokasi dengan keluaran DN40", BB); put(ws, r, 6, f"=SUM(F5:F{last})", BB, RP, TINT); PER_LOK = f"'Biaya dasar'!F{r}"; r += 2

put(ws, r, 1, "Lokasi", BB); put(ws, r, 2, "Meter keluar", BB); put(ws, r, 3, "Meter masuk", BB); put(ws, r, 4, "Penyesuaian (Rp)", BB); put(ws, r, 6, "Biaya lokasi (Rp)", BB); r += 1
lok = [("Mander, Tambakboyo", "DN40", "DN20", -50000), ("Bader, Jatirogo", "DN40", "DN25", 0), ("Tuwiri Kulon, Merakurak", "DN15", "DN25", -1100000), ("Kabluan, Bangilan", "DN40", "DN20", -50000), ("Ngulahan, Tambakboyo", "DN40", "DN25", 0)]
r0 = r
for n, mk, mm, adj in lok:
    put(ws, r, 1, n); put(ws, r, 2, mk); put(ws, r, 3, mm); c = put(ws, r, 4, adj, BLUE, RP); put(ws, r, 6, f"={PER_LOK}+D{r}", B, RP); r += 1
ws.cell(row=r0, column=4).comment = Comment("DN20 lebih murah Rp50.000 dari DN25; DN15 + strainer kecil lebih murah Rp1.100.000 dari DN40. Angka masukan dari pengusul.", "Ratna")
put(ws, r, 1, "Lima lokasi", BB); put(ws, r, 6, f"=SUM(F{r0}:F{r-1})", BB, RP, TINT); LIMA = f"'Biaya dasar'!F{r}"; r += 1
put(ws, r, 1, "Perjalanan tim Surabaya–Tuban", B); put(ws, r, 2, "3 kunjungan × 5 desa", B); put(ws, r, 4, 3, BLUE); put(ws, r, 5, 1500000, BLUE, RP); put(ws, r, 6, f"=D{r}*E{r}", B, RP); PERJ = f"'Biaya dasar'!F{r}"; r += 1
put(ws, r, 1, "Cadangan komponen (mitigasi)", B); put(ws, r, 2, "1 sensor level, 1 modul 4G, 1 ESP32", B); put(ws, r, 6, "=E8+E11+E10", B, RP); CAD = f"'Biaya dasar'!F{r}"; r += 1
put(ws, r, 1, "BIAYA DASAR (cost estimate)", BB); put(ws, r, 6, f"={LIMA}+{PERJ}+{CAD}", BB, RP, TINT); DASAR = f"'Biaya dasar'!F{r}"
for col, w in zip("ABCDEF", [6, 40, 36, 8, 18, 18]): ws.column_dimensions[col].width = w

# ---------------- Sheet 2: Daftar risiko ----------------
rs = wb.create_sheet("Daftar risiko")
rs["A1"] = "Daftar risiko dan cadangan risiko (contingency reserve) — nilai harapan = peluang × dampak biaya"; rs["A1"].font = Font(name=F, bold=True, size=13)
rs["A2"] = "Ubah sel biru (peluang dan dampak) sesuai kondisi lapangan. Dampak = biaya tambahan bila risiko terjadi, dihitung dari rumus di kolom Cara menghitung dampak."; rs["A2"].font = Font(name=F, italic=True, size=9)
hdr(rs, 4, ["ID", "Risiko", "Pemicu / sumber", "Cara menghitung dampak", "Peluang", "Dampak biaya (Rp)", "Nilai harapan (Rp)", "Dampak jadwal (hari)", "Mitigasi (dicantumkan dalam biaya dasar atau rencana)"])
risks = [
 ("R1", "Komponen terlambat datang (water meter, sensor level, modul 4G)", "Stok penjual kosong, pengiriman impor 2–4 minggu", "Satu kunjungan tambahan untuk memasang komponen yang tertinggal: 1 × biaya kunjungan", 0.30, 1500000, 14, "Pesan di minggu ke-2, pilih penjual stok siap, cadangan 1 unit sensor dan modul"),
 ("R2", "Komponen rusak saat datang atau saat uji (DOA)", "Sensor level, modul 4G", "Harga pengganti + ongkir kilat: 500.000 + 100.000", 0.25, 600000, 5, "Uji semua komponen di lab sebelum berangkat; cadangan 1 unit"),
 ("R3", "Tukang pipa tidak hadir saat pemasangan, pengganti lokal lebih mahal", "Pelanggaran kehadiran, 2 kunjungan pemasangan", "Selisih tarif pengganti (450.000 − 250.000) + setengah hari tim menunggu (200.000), per kunjungan; 2 kunjungan", 0.40, 800000, 1, "Kontrak dengan uang muka, dua kontak tukang per desa, konfirmasi H-1"),
 ("R4", "Anggota tim mahasiswa tidak hadir atau mengundurkan diri", "Jadwal kuliah, sakit, pelanggaran kehadiran", "Honor asisten lapangan pengganti 1 kunjungan 500.000 + lembur perakitan 250.000", 0.30, 750000, 3, "Tim 4 orang untuk kebutuhan 3; jadwal pemasangan di luar minggu ujian"),
 ("R5", "Sinyal IM3 lemah di salah satu rumah pompa", "Lokasi di lembah atau jauh dari BTS", "Antena eksternal + kabel 200.000 + kunjungan ulang sebagian 500.000", 0.40, 700000, 7, "Cek sinyal saat survei; bawa antena cadangan saat pemasangan"),
 ("R6", "Ukuran pipa di lapangan berbeda dari catatan survei", "Catatan survei keliru, pipa sudah diganti", "Beli ulang satu meter + modul pulsa 1.200.000", 0.20, 1200000, 10, "Ukur ulang dan foto pipa saat survei; beli meter setelah survei"),
 ("R7", "Kenaikan harga komponen atau ongkos kirim", "Kurs, harga marketplace berubah", "10% dari 60% komponen yang belum terbeli: 0,1 × 0,6 × 20.000.000", 0.40, 1200000, 0, "Beli komponen utama di minggu pertama setelah dana cair"),
 ("R8", "Akses lokasi tertunda (izin desa, hujan, jalan rusak)", "Musim hujan November–Januari di Tuban", "Satu kunjungan ulang 1.500.000", 0.25, 1500000, 7, "Koordinasi dengan kepala desa sebelum berangkat; pemasangan 3 lokasi per kunjungan"),
 ("R9", "Kalibrasi level gagal, sensor tidak cocok dengan nipel toren", "Ulir nipel berbeda, toren tanpa nipel bawah", "Fitting adaptor 150.000 + setengah hari kerja 350.000", 0.30, 500000, 3, "Bawa set adaptor ulir ¼, ½, ¾ inci; foto nipel saat survei"),
]
r = 5
for rid, nm, src, how, p, imp, days, mit in risks:
    put(rs, r, 1, rid); put(rs, r, 2, nm); put(rs, r, 3, src); put(rs, r, 4, how)
    put(rs, r, 5, p, BLUE, "0%"); put(rs, r, 6, imp, BLUE, RP); put(rs, r, 7, f"=E{r}*F{r}", B, RP); put(rs, r, 8, days, BLUE); put(rs, r, 9, mit); r += 1
rl = r - 1
put(rs, r, 2, "CADANGAN RISIKO (jumlah nilai harapan)", BB); put(rs, r, 7, f"=SUM(G5:G{rl})", BB, RP, TINT); CONT = f"'Daftar risiko'!G{r}"; r += 1
put(rs, r, 2, "Dampak terburuk bila semua terjadi", B); put(rs, r, 6, f"=SUM(F5:F{rl})", B, RP); r += 1
put(rs, r, 2, "Buffer jadwal yang disarankan (hari) = jumlah dampak jadwal × peluang", B); put(rs, r, 8, f"=SUMPRODUCT(E5:E{rl},H5:H{rl})", B, "0"); r += 1
for col, w in zip("ABCDEFGHI", [5, 38, 30, 44, 9, 16, 16, 11, 44]): rs.column_dimensions[col].width = w
rs.row_dimensions[4].height = 32

# ---------------- Sheet 3: Anggaran total ----------------
ts = wb.create_sheet("Anggaran total")
ts["A1"] = "Anggaran total dengan konsep manajemen proyek (PMBOK: cost baseline + management reserve)"; ts["A1"].font = Font(name=F, bold=True, size=13)
rows = [
 ("Biaya dasar (estimasi bottom-up, termasuk cadangan komponen)", f"={DASAR}", "Biaya dasar"),
 ("Cadangan risiko (contingency reserve) = Σ peluang × dampak dari daftar risiko", f"={CONT}", "Daftar risiko"),
 ("Cadangan risiko sebagai persen biaya dasar", "=B6/B5", None),
 ("COST BASELINE = biaya dasar + cadangan risiko", "=B5+B6", None),
 ("Persen cadangan manajemen (risiko yang belum teridentifikasi)", 0.05, "masukan"),
 ("Cadangan manajemen (management reserve) = persen × cost baseline", "=B9*B8", None),
 ("ANGGARAN TOTAL YANG DIAJUKAN = cost baseline + cadangan manajemen", "=B8+B10", None),
 ("Pembulatan ke atas (Rp100.000)", "=CEILING(B11,100000)", None),
]
hdr(ts, 4, ["Komponen anggaran", "Nilai", "Sumber"])
r = 5
for lab, val, src in rows:
    put(ts, r, 1, lab, BB if lab.isupper() or lab.startswith(("COST", "ANGGARAN")) else B)
    c = put(ts, r, 2, val, BLUE if src == "masukan" else B, "0.0%" if "persen" in lab.lower() and "Cadangan manajemen (" not in lab else RP, YEL if src == "masukan" else (TINT if lab.startswith(("COST", "ANGGARAN")) else None))
    put(ts, r, 3, src or "rumus"); r += 1
ts.cell(row=7, column=2).number_format = "0.0%"
ts.cell(row=9, column=2).comment = Comment("5% dipakai untuk proyek kecil dengan risiko yang sudah dirinci; 10% bila lokasi belum disurvei. Masukan dari pengusul.", "Ratna")
r += 1
ts.cell(row=r, column=1, value="Aturan pemakaian cadangan").font = BB; r += 1
for t in ["Cadangan risiko dipakai ketua tim untuk risiko yang ada di daftar, dicatat tanggal, risiko, dan nominalnya di log proyek.",
          "Cadangan manajemen hanya dipakai untuk kejadian di luar daftar risiko, dengan persetujuan pengusul/LPPM.",
          "Sisa cadangan yang tidak terpakai dikembalikan atau dialihkan ke tahap lanjutan (meteran per rumah) dengan persetujuan pemberi dana.",
          "Setiap akhir minggu, biaya aktual dibandingkan dengan cost baseline (earned value sederhana: rencana vs realisasi) untuk melihat apakah cadangan mulai terpakai."]:
    ts.cell(row=r, column=1, value="• " + t).font = B; ts.cell(row=r, column=1).alignment = Alignment(wrap_text=True); r += 1
ts.column_dimensions["A"].width = 80; ts.column_dimensions["B"].width = 20; ts.column_dimensions["C"].width = 16

wb.save("RAB Monitoring Toren Desa dengan Cadangan Risiko.xlsx"); print("saved")
